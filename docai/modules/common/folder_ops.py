"""Direct file operations in the working folder via chat — create a new
file, delete a file (with undo), rename a file, or open + edit a named
file — per the design flow at
`D:\\Tools\\DocAI\\Plan\\V4\\Folder\\screens_folder`.

These are LOCAL operations, no AI call for classification (Vietnamese
regex) — unlike `doc_set.py` (analyze/compare/merge MULTIPLE files).
Recognition here takes priority; if nothing matches, fall through to
`doc_set.process_document_set()`.
"""
import os
import re
import shutil
import time
import unicodedata
from dataclasses import dataclass

from ...account import api_client

from ...logging_config import get_logger, log_call
from ...strings import FolderOps as S

logger = get_logger(__name__)

_TRASH_DIRNAME = ".docai_trash"

_EXT_TO_TYPE = {
    ".docx": "word", ".doc": "word",
    ".xlsx": "excel", ".xls": "excel",
    ".pptx": "ppt", ".ppt": "ppt",
}
_TYPE_TO_EXT = {"word": ".docx", "excel": ".xlsx", "ppt": ".pptx"}
_TYPE_WORDS = {
    "excel": "excel", "bảng tính": "excel",
    "powerpoint": "ppt", "trình chiếu": "ppt", "trình bày": "ppt", "ppt": "ppt",
    "word": "word", "văn bản": "word",
}


class FolderOpError(api_client.ApiError):
    @log_call
    def __init__(self, message: str):
        super().__init__(message, "folder_op_failed")


_CAMEL_RE = re.compile(r"(?<=[a-z0-9])(?=[A-Z])")


@log_call
def _norm(text: str) -> str:
    """Split camelCase (e.g. "HopDong" → "Hop Dong") before normalizing —
    file names are often concatenated this way, no space between words."""
    text = _CAMEL_RE.sub(" ", text)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-zA-Z0-9]+", " ", text.lower()).strip()


@log_call
def _name_mentioned(text_words: set[str], path: str, threshold: float = 0.75) -> bool:
    """Word-overlap, higher threshold than `doc_set.py` because it's used to
    infer the file for delete/rename/open — needs more certainty than the
    "include in analysis" bar."""
    stem_words = {w for w in _norm(os.path.splitext(os.path.basename(path))[0]).split() if len(w) >= 2}
    if not stem_words:
        return False
    return len(stem_words & text_words) / len(stem_words) >= threshold


_QUOTED_RE = re.compile(r'["\'“”](.+?)["\'“”]')


@log_call
def _quoted_names(text: str) -> list[str]:
    return [m.group(1).strip() for m in _QUOTED_RE.finditer(text or "")]


@log_call
def _resolve_by_name(target_name: str | None, text: str, folder_files: list[str]) -> str | None:
    """Prefer an exact match on a name extracted from quotes; if none/no
    match, try recognizing the file name mentioned in the sentence (only
    accepted if it's a UNIQUE match)."""
    if target_name:
        target_norm = _norm(os.path.splitext(target_name)[0])
        for path in folder_files:
            if _norm(os.path.splitext(os.path.basename(path))[0]) == target_norm:
                return path
    text_words = {w for w in _norm(text).split() if len(w) >= 2}
    matches = [path for path in folder_files if _name_mentioned(text_words, path)]
    return matches[0] if len(matches) == 1 else None


# ── Command recognition (local) ──────────────────────────────────────────────

_CREATE_RE = re.compile(r"\btạo\s+(file|tệp)\b", re.IGNORECASE)
_DELETE_RE = re.compile(r"\bxóa\s+(file|tệp)\b", re.IGNORECASE)
_RENAME_RE = re.compile(r"\bđổi\s+tên\b", re.IGNORECASE)
_OPEN_RE = re.compile(r"\bmở\b", re.IGNORECASE)
_RENAME_SPLIT_RE = re.compile(r"(.+?)\s+thành\s+(.+)", re.IGNORECASE)


@dataclass
class FileMgmtIntent:
    kind: str                       # "create" | "delete" | "rename" | "open"
    name: str | None = None         # create: new file name
    ftype: str | None = None        # create: "word"|"excel"|"ppt"
    target_path: str | None = None  # delete/rename/open: matched file in the folder
    new_name: str | None = None     # rename: the new name
    remainder: str = ""             # open: the remaining edit command after the file name
    ambiguous: bool = False         # missing info — ask back instead of guessing


@log_call
def _parse_create(text: str) -> FileMgmtIntent:
    names = _quoted_names(text)
    name = names[0] if names else None
    ftype = None
    lower = text.lower()
    for word, mapped in _TYPE_WORDS.items():
        if word in lower:
            ftype = mapped
            break
    if name:
        ext = os.path.splitext(name)[1].lower()
        if ext in _EXT_TO_TYPE:
            ftype = ftype or _EXT_TO_TYPE[ext]
        elif ftype:
            name = name + _TYPE_TO_EXT[ftype]
    return FileMgmtIntent(kind="create", name=name, ftype=ftype,
                          ambiguous=not (name and ftype))


@log_call
def _parse_delete(text: str, folder_files: list[str]) -> FileMgmtIntent:
    names = _quoted_names(text)
    target = names[0] if names else None
    path = _resolve_by_name(target, text, folder_files)
    return FileMgmtIntent(kind="delete", target_path=path, ambiguous=path is None)


@log_call
def _parse_rename(text: str, folder_files: list[str]) -> FileMgmtIntent:
    names = _quoted_names(text)
    if len(names) >= 2:
        old_name, new_name = names[0], names[1]
    else:
        match = _RENAME_SPLIT_RE.search(text)
        old_name, new_name = (match.group(1).strip(), match.group(2).strip()) if match else (None, None)
    path = _resolve_by_name(old_name, text, folder_files) if old_name else None
    return FileMgmtIntent(kind="rename", target_path=path, new_name=new_name,
                          ambiguous=not (path and new_name))


@log_call
def _parse_open(text: str, folder_files: list[str]) -> FileMgmtIntent | None:
    names = _quoted_names(text)
    target = names[0] if names else None
    path = _resolve_by_name(target, text, folder_files)
    if path is None:
        return FileMgmtIntent(kind="open", ambiguous=True)

    remainder = text
    if target:
        idx = text.find(target)
        if idx != -1:
            remainder = text[idx + len(target):]
    remainder = re.sub(
        r'^[\s"\'“”]*(trong\s+(thư mục( này| làm việc)?|đây))?[\s,]*(và)?\s*',
        "", remainder, flags=re.IGNORECASE).strip()
    return FileMgmtIntent(kind="open", target_path=path, remainder=remainder)


@log_call
def detect_file_mgmt_intent(text: str, folder_files: list[str]) -> FileMgmtIntent | None:
    """None if the chat message doesn't match any create/delete/rename/open
    command — the caller should fall through to other handling
    (`doc_set.process_document_set`)."""
    if _CREATE_RE.search(text):
        return _parse_create(text)
    if _DELETE_RE.search(text):
        return _parse_delete(text, folder_files)
    if _RENAME_RE.search(text):
        return _parse_rename(text, folder_files)
    if _OPEN_RE.search(text):
        return _parse_open(text, folder_files)
    return None


# ── Execution (real I/O) ──────────────────────────────────────────────────────

@log_call
def create_empty_file(folder: str, name: str, ftype: str) -> str:
    path = os.path.join(folder, name)
    if os.path.exists(path):
        raise FolderOpError(S.FILE_ALREADY_EXISTS.format(name=name))
    try:
        if ftype == "word":
            from docx import Document
            Document().save(path)
        elif ftype == "excel":
            import openpyxl
            openpyxl.Workbook().save(path)
        elif ftype == "ppt":
            from pptx import Presentation
            Presentation().save(path)
        else:
            raise FolderOpError(S.FILE_TYPE_UNSUPPORTED.format(ftype=ftype))
    except FolderOpError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise FolderOpError(S.CREATE_FAILED.format(error=exc)) from exc
    return path


@log_call
def delete_file_with_undo(path: str) -> str:
    """Move into `.docai_trash` alongside the folder instead of deleting for
    real — returns the temp path so `restore_file()` can put it back in its
    original location."""
    folder = os.path.dirname(path)
    trash_dir = os.path.join(folder, _TRASH_DIRNAME)
    os.makedirs(trash_dir, exist_ok=True)
    dest = os.path.join(trash_dir, os.path.basename(path))
    if os.path.exists(dest):
        base, ext = os.path.splitext(dest)
        dest = f"{base}_{int(time.time())}{ext}"
    try:
        shutil.move(path, dest)
    except OSError as exc:
        raise FolderOpError(S.DELETE_FAILED.format(error=exc)) from exc
    return dest


@log_call
def restore_file(trashed_path: str, original_path: str) -> None:
    if os.path.exists(original_path):
        raise FolderOpError(S.RESTORE_NAME_CONFLICT)
    try:
        shutil.move(trashed_path, original_path)
    except OSError as exc:
        raise FolderOpError(S.RESTORE_FAILED.format(error=exc)) from exc


@log_call
def rename_file(path: str, new_name: str) -> str:
    directory = os.path.dirname(path)
    old_ext = os.path.splitext(path)[1]
    if not os.path.splitext(new_name)[1]:
        new_name += old_ext
    new_path = os.path.join(directory, new_name)
    if os.path.exists(new_path):
        raise FolderOpError(S.RENAME_NAME_CONFLICT.format(name=new_name))
    try:
        os.rename(path, new_path)
    except OSError as exc:
        raise FolderOpError(S.RENAME_FAILED.format(error=exc)) from exc
    return new_path
